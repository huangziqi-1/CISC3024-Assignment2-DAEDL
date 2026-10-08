#!/usr/bin/env python3
"""Generate the final report source and editable Word document from validated data."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SUMMARY = ROOT / "results/final_summary"
FIGURES = ROOT / "figures/final"
REPORT = ROOT / "report"
EQUATIONS = REPORT / "equations"

def load_results() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    aggregate = pd.read_csv(SUMMARY / "aggregate_metrics.csv").set_index(["model", "metric"])
    per_seed = pd.read_csv(SUMMARY / "per_seed_metrics.csv")
    density = pd.read_csv(SUMMARY / "density_diagnostics_by_seed.csv")
    checks = json.loads((SUMMARY / "consistency_checks.json").read_text())
    if checks.get("status") != "PASS":
        raise AssertionError("final report requires a passing Phase 3B audit")
    if set(per_seed.seed) != {3024, 3025, 3026} or set(per_seed.model) != {"softmax", "daedl"}:
        raise AssertionError("unexpected seed/model set")
    return aggregate, per_seed, density


def value(aggregate: pd.DataFrame, model: str, metric: str) -> str:
    row = aggregate.loc[(model, metric)]
    return f"{row['mean']:.4f} ± {row['sample_sd']:.4f}"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9D9D9", size: str = "6") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, amount in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(amount))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "17365D")
        set_cell_border(cell)
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column, item in enumerate(values):
            cells[column].text = str(item)
            cells[column].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cells[column])
            set_cell_margins(cells[column])
            if row_index % 2:
                set_cell_shading(cells[column], "F2F6FA")
            for paragraph in cells[column].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    return table


def add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, end])


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relation = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def remove_paragraph_border(paragraph_or_style) -> None:
    element = paragraph_or_style._p if hasattr(paragraph_or_style, "_p") else paragraph_or_style.element
    properties = element.get_or_add_pPr()
    border = properties.find(qn("w:pBdr"))
    if border is not None:
        properties.remove(border)


def add_caption(document: Document, label: str, text: str):
    paragraph = document.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"{label}. ")
    run.bold = True
    paragraph.add_run(text)
    keep_with_next(paragraph)
    return paragraph


def add_figure(document: Document, path: Path, width: float):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    return paragraph


def add_equation(document: Document, name: str, width: float = 6.2):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.add_run().add_picture(str(EQUATIONS / name), width=Inches(width))


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = ""
    section.first_page_footer.paragraphs[0].text = ""

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size, before, after in (
        ("Title", 22, 0, 12), ("Heading 1", 15, 10, 5),
        ("Heading 2", 12.5, 8, 4), ("Heading 3", 11, 6, 3),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    remove_paragraph_border(styles["Title"])
    styles["Caption"].font.name = "Aptos"
    styles["Caption"].font.size = Pt(9)
    styles["Caption"].font.italic = False
    styles["Caption"].font.color.rgb = RGBColor(45, 45, 45)
    styles["Caption"].paragraph_format.space_after = Pt(4)
    if "References" not in styles:
        reference_style = styles.add_style("References", WD_STYLE_TYPE.PARAGRAPH)
        reference_style.base_style = styles["Normal"]
        reference_style.font.name = "Aptos"
        reference_style.font.size = Pt(9)
        reference_style.paragraph_format.left_indent = Inches(0.22)
        reference_style.paragraph_format.first_line_indent = Inches(-0.22)
        reference_style.paragraph_format.space_after = Pt(4)

    for page_section in document.sections:
        header = page_section.header.paragraphs[0]
        header.text = "CISC3024 Pattern Recognition  |  AI Assignment #2"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(90, 90, 90)
        footer = page_section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Page ")
        run.font.size = Pt(8)
        add_field(footer.add_run(), "PAGE")


def paragraph(document: Document, text: str, bold_lead: str | None = None):
    p = document.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def bullet(document: Document, text: str):
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def build_markdown(aggregate: pd.DataFrame, density: pd.DataFrame, source_url: str | None) -> str:
    s_acc, d_acc = value(aggregate, "softmax", "accuracy"), value(aggregate, "daedl", "accuracy")
    s_f1, d_f1 = value(aggregate, "softmax", "macro_f1"), value(aggregate, "daedl", "macro_f1")
    s_nll, d_nll = value(aggregate, "softmax", "nll"), value(aggregate, "daedl", "nll")
    s_ece, d_ece = value(aggregate, "softmax", "ece_15"), value(aggregate, "daedl", "ece_15")
    s_brier, d_brier = value(aggregate, "softmax", "brier"), value(aggregate, "daedl", "brier")
    s_err, d_err = value(aggregate, "softmax", "error_auroc"), value(aggregate, "daedl", "error_auroc")
    s_errp, d_errp = value(aggregate, "softmax", "error_aupr"), value(aggregate, "daedl", "error_aupr")
    s_ood, d_ood = value(aggregate, "softmax", "ood_auroc"), value(aggregate, "daedl", "ood_auroc")
    s_oodp, d_oodp = value(aggregate, "softmax", "ood_aupr"), value(aggregate, "daedl", "ood_aupr")
    source_section = (
        "\n## 11 Source Code\n\n"
        "The complete source code, experiment scripts, configurations, and reproducibility files "
        f"for this assignment are available at:\n\n**{source_url}**\n"
        if source_url
        else ""
    )
    references_heading = "## 12 References" if source_url else "## References"
    return fr"""# Density-Aware Evidential Deep Learning for CIFAR-10 Classification and Uncertainty Estimation

**CISC3024 Pattern Recognition**

**AI Assignment #2**

Ziqi Huang

October 2026

## 1 Introduction

Image classifiers produce probability-like scores, but a high maximum score does not by itself show that the model knows when it may be wrong. This assignment used Codex to identify, implement, and evaluate a recent probability-based recognition method. The selected method was Density Aware Evidential Deep Learning (DAEDL), published at ICML 2024 [1]. The research question was: how does DAEDL compare with a conventional Softmax classifier in classification performance, calibration, prediction-error uncertainty, and out-of-distribution uncertainty in an assignment-scale CIFAR-10 experiment?

The three-seed results show different strengths across the evaluation metrics. Softmax achieved higher accuracy ({s_acc} versus {d_acc}) and better NLL and Brier score. DAEDL achieved lower ECE in every seed ({d_ece} versus {s_ece}). Both methods assigned greater entropy to mistakes, and Softmax had a small error-AUROC advantage. The SVHN OOD ranking changed across seeds, so neither method was consistently better.

## 2 AI-Assisted Algorithm Search and Selection

Codex searched for recent classification and recognition methods grounded in probability or Bayesian theory. Codex first compared DAEDL with conventional evidential deep learning. The candidates were screened by publication date, probabilistic basis, implementation availability, computational cost, and suitability for a CIFAR-10 comparison. DAEDL was selected because it predicts a Dirichlet distribution, models density in learned feature space, and supports classification, calibration, error-detection, and OOD evaluation [1].

The search then checked the official PMLR record, the paper, and the authors' repository. Codex used the repository as a method reference and created a separate implementation adapted to this assignment. The repository audit found CUDA assumptions, incomplete environment information, missing released checkpoints, and reproducibility gaps. Conventional evidential deep learning [2] and calibration work [3] supplied the background needed to interpret the method and metrics.

## 3 Algorithm Description

### 3.1 Softmax classification

For class logits $z_k(x)$, Softmax forms one normalized categorical probability vector for each input:

$$p_{{soft}}(y=k\mid x)=\frac{{e^{{z_k(x)}}}}{{\sum_j e^{{z_j(x)}}}}.$$

The largest probability determines the predicted class. Softmax does not directly represent uncertainty over the probability vector itself, which motivates the Dirichlet formulation used by evidential classifiers.

### 3.2 Evidential classification

Evidential classification predicts positive concentrations $\boldsymbol{{\alpha}}(x)$ for a Dirichlet distribution over class-probability vectors $\boldsymbol{{\pi}}$. Its predictive categorical probability is the Dirichlet mean, $\mathbb{{E}}[\pi_k\mid x]=\alpha_k/S$, where $S=\sum_j\alpha_j$. A larger total concentration $S$ places more Dirichlet mass around the mean probability vector; a smaller value represents weaker evidence. This project used normalized predictive entropy as the common uncertainty score so that Softmax and DAEDL were compared on the same quantity.

### 3.3 DAEDL training and density-aware prediction

DAEDL uses $\alpha_k=e^{{z_k}}$ during training. Normalizing these concentrations gives the same class probabilities as Softmax, while the loss adds the expected squared-error and predictive-variance terms of the Dirichlet plus a KL penalty for misleading evidence. After training, the implementation fits class-conditional full-covariance Gaussian distributions to the 128-dimensional penultimate embeddings. The class mixture defines a feature density, which is min-max normalized against training log densities and clipped to $[0,1]$.

At inference, the density score $s(x)$ scales the logits before concentrations are formed: $\alpha_k(x)=\exp(s(x)z_k(x))$. An embedding far from the fitted training distribution receives a small density score. This moves the concentrations toward one and the expected class probabilities toward a uniform distribution. Spectral normalization is applied to both compared models so that the backbone treatment is matched; for DAEDL it also supports a smoother feature geometry for density estimation.

![Figure 1 DAEDL concept](../figures/final/figure1_daedl_concept.png)

## 4 AI-Assisted Implementation

### 4.1 Project and environment setup

Codex created an isolated Python 3.11.5 environment with PyTorch 2.14.1 and torchvision 0.29.1 on an Apple M5 MacBook Air with 16 GB unified memory. The project code covers data loading, model construction, Softmax and DAEDL training, density fitting, evaluation, artifact logging, and report figures. Neural-network training used MPS. Selected checkpoints were evaluated on CPU, and DAEDL density estimation used CPU float64 because the MPS backend cannot store float64 tensors [6].

The pilot also found that a reloaded spectral-normalized checkpoint could produce non-finite outputs on the observed MPS evaluation path while the same weights were finite on CPU. The final protocol therefore used the same CPU checkpoint-evaluation policy for both models.

### 4.2 Data pipeline

Each CIFAR-10 example retained its original sample index so that splits and predictions could be audited. The 50,000 official training images were divided into a stratified 45,000-image training set and 5,000-image validation set. Training used random 32-by-32 crops with four-pixel padding, random horizontal flips, tensor conversion, and channel normalization with mean $(0.4914, 0.4822, 0.4465)$ and standard deviation $(0.2023, 0.1994, 0.2010)$. Validation, test, and density fitting used only tensor conversion and the same normalization.

The code saved the train, validation, and test indices for every run. Python, NumPy, PyTorch, and the data-loader generator were seeded. Within a seed, Softmax and DAEDL used identical indices; across seeds, the split and training randomness changed according to the recorded seed.

### 4.3 Model implementation

Both methods used the same CIFAR-adapted ResNet-18. It has a 3-by-3 stride-one stem, base width 32, four residual stages, adaptive average pooling, a 128-dimensional ReLU embedding, and a ten-class linear classifier. The forward method exposes both logits and embeddings because DAEDL needs the embeddings after training.

Spectral normalization was applied to every convolution, learned shortcut projection, the embedding layer, and the classifier for both models. Sharing the feature extractor, data pipeline, and training budget keeps the comparison focused on the prediction and uncertainty mechanisms. Softmax uses cross-entropy. DAEDL uses its evidential objective, followed by feature-density fitting and density-aware inference.

### 4.4 DAEDL density estimation

Density estimation begins after neural-network training. The selected checkpoint is loaded on CPU, and clean embeddings are extracted from all 45,000 training images. The implementation then fits one full-covariance Gaussian per class, using empirical class proportions as mixture weights. The minimum and maximum training log densities define the clipped score $s(x)$ used at inference.

All covariance calculations use CPU float64. Each covariance is symmetrized, and scale-relative diagonal jitter with an absolute floor of $10^{{-8}}$ is increased until Cholesky factorization succeeds. The run artifacts record global and per-class ranks, constant dimensions, jitter, the training density range, and non-finite fractions. The fitted GDA state and per-example training density values are also saved.

### 4.5 Staged implementation workflow

The implementation was validated in stages. A two-epoch pilot used 5,000 training, 1,000 validation, and 1,000 test images to check finite losses and gradients, checkpoint reloads, GDA fitting, metrics, and artifact output. The first full experiment then trained both models on seed 3024 and evaluated classification, calibration, prediction errors, SVHN OOD detection, and density diagnostics. After those checks passed, the same configuration was then used for seeds 3025 and 3026.

### 4.6 Validation, testing, and reproducibility

The 16 automated tests cover four groups. Probability and loss tests check positive concentrations, normalized probabilities, and finite DAEDL losses and gradients. Metric tests check NLL, Brier score, ECE, entropy, and AUROC/AUPR direction. Density tests cover shapes, CPU float64 state, rank-deficient covariance handling, and finite scores. Reproducibility tests cover stratified disjoint splits, exposed model outputs, and spectral-normalization coverage.

Every final run saved its configuration, environment, split indices, progress log, training history, best-validation and final checkpoints, per-example CIFAR-10 and SVHN predictions, metrics, runtime, and DAEDL density diagnostics where applicable. Aggregate tables were recomputed from the saved predictions. The final six-run audit verified matched within-seed splits, checkpoint selection, probability normalization, finite outputs, positive DAEDL concentrations, score direction, density stability, and sample-SD aggregation.

## 5 Experimental Setup

### 5.1 Datasets and split

CIFAR-10 contains ten natural-object classes with 50,000 official training images and 10,000 official test images [4]. The training portion was divided into a stratified 45,000-image training set and 5,000-image validation set. Validation NLL selected checkpoints, while the official test set remained outside model selection.

SVHN supplied 10,000 deterministically selected test images for OOD evaluation [5]. CIFAR-10 contains natural objects; SVHN contains street-view house-number digits. Their visual domains and semantic classes differ, making SVHN a useful OOD dataset for this controlled comparison without implying that it is the only possible benchmark.

### 5.2 Fair comparison design

Softmax and DAEDL shared the same backbone, spectral-normalization configuration, split, training augmentation, batch size, optimizer, epoch budget, learning-rate schedule, and checkpoint criterion within each seed. These controls reduce differences caused by the feature extractor or training resources. DAEDL necessarily differs in its evidential objective, post-training density model, and density-aware inference.

### 5.3 Training protocol

Each run trained for 30 epochs with batch size 64, AdamW, an initial learning rate of $10^{{-3}}$, weight decay $5\times10^{{-4}}$, and cosine decay to $10^{{-5}}$. Both models selected the checkpoint with minimum validation NLL. This provides a common probability-based criterion even though their training objectives differ. The best-validation checkpoint and the final-epoch checkpoint were both saved, and the test set was never used for selection.

| Setting | Value |
|---|---|
| Models | Softmax and DAEDL |
| Backbone | Spectrally normalized CIFAR ResNet-18, base width 32 |
| Embedding | 128 dimensions |
| Split | 45,000 train / 5,000 validation / 10,000 test |
| Seeds | 3024, 3025, 3026 |
| Optimizer | AdamW, LR $10^{{-3}}$, weight decay $5\times10^{{-4}}$ |
| Schedule | Cosine decay to $10^{{-5}}$ over 30 epochs |
| Batch size | 64 |
| Selection | Minimum validation NLL |
| Hardware | Apple M5 MacBook Air, 16 GB unified memory |
| Devices | MPS training; CPU evaluation and float64 GDA |
| OOD protocol | CIFAR-10 ID, SVHN OOD, normalized entropy, OOD positive |

### 5.4 Three-seed replication

Seeds 3024, 3025, and 3026 each defined a complete model run. The seed controlled the stratified train-validation split, Python and NumPy random states, PyTorch initialization and training randomness, and the data-loader generator. Softmax and DAEDL shared indices within each seed. Results are reported as mean ± sample SD across the three runs, not across individual images. Three seeds were used to check whether the comparison remained consistent under different data splits and training randomness.

### 5.5 Classification evaluation

Accuracy is the fraction of correct CIFAR-10 predictions. Macro F1 averages class-level F1 scores with equal weight for each class. CIFAR-10 is balanced, but Macro F1 still checks whether performance is distributed across classes instead of relying only on total accuracy.

### 5.6 Calibration evaluation

Negative log-likelihood uses the probability assigned to the true class and gives a large penalty to confident mistakes. ECE divides predictions into 15 equal-width confidence bins and computes the count-weighted difference between mean confidence and observed accuracy in each bin. The multiclass Brier score is the mean squared difference between the predicted probability vector and the one-hot target. Using all three metrics gives a broader view of probability quality than relying on a single calibration measure.

### 5.7 Prediction-error uncertainty

This experiment asks whether uncertainty can identify predictions that are likely to be wrong. Correct and incorrect CIFAR-10 predictions are separated using the ground-truth labels. Normalized predictive entropy lies in $[0,1]$; larger values indicate a more diffuse class distribution. Prediction error is the positive class. AUROC measures ranking across thresholds, and AUPR summarizes precision-recall performance for the error class.

### 5.8 OOD evaluation

The OOD experiment asks whether the models assign greater uncertainty to inputs outside the CIFAR-10 distribution. CIFAR-10 is labelled ID, SVHN is labelled OOD, and OOD is the positive class. Normalized predictive entropy is the common score for both models, which avoids comparing method-specific scales. DAEDL's secondary score, $1-s(x)$, is evaluated separately as density-derived uncertainty.

### 5.9 Numerical stability and device policy

Training used MPS, while loaded-checkpoint evaluation and feature extraction used CPU for both models. DAEDL covariance and density calculations used CPU float64, adaptive diagonal jitter, and Cholesky checks. Evaluation also checked finite outputs, probability sums, positive concentrations, density bounds, and score orientation before aggregation.

### 5.10 Experiment questions

| Experiment | Main question | Metrics |
|---|---|---|
| Classification | How accurately does the model classify CIFAR-10? | Accuracy, Macro F1 |
| Calibration | Do predicted probabilities match empirical correctness? | NLL, ECE, Brier |
| Error uncertainty | Can uncertainty identify incorrect predictions? | Entropy, AUROC, AUPR |
| OOD uncertainty | Can uncertainty distinguish CIFAR-10 from SVHN? | Entropy, AUROC, AUPR |

## 6 Results

### 6.1 Classification and calibration

| Metric | Softmax | DAEDL |
|---|---:|---:|
| Accuracy | {s_acc} | {d_acc} |
| Macro F1 | {s_f1} | {d_f1} |
| NLL | {s_nll} | {d_nll} |
| ECE-15 | {s_ece} | {d_ece} |
| Brier score | {s_brier} | {d_brier} |

Softmax was about one percentage point more accurate in every seed. DAEDL's lower ECE also repeated in every seed, but NLL and Brier score favored Softmax. The calibration metrics therefore do not support one universal ranking. ECE measures binned agreement between confidence and observed accuracy, whereas NLL and Brier score evaluate the complete predictive distribution as proper scoring rules.

![Figure 2 three-seed headline metrics](../figures/final/figure2_three_seed_headline.png)

The plot shows seed 3024 as an example. Bins with fewer than 50 test examples are suppressed for readability, while ECE is calculated from all 15 equal-width bins.

![Figure 3 reliability](../figures/final/figure3_reliability_seed3024.png)

### 6.2 Error and OOD uncertainty

| Metric | Softmax | DAEDL |
|---|---:|---:|
| Error AUROC | {s_err} | {d_err} |
| Error AUPR | {s_errp} | {d_errp} |
| SVHN OOD AUROC | {s_ood} | {d_ood} |
| SVHN OOD AUPR | {s_oodp} | {d_oodp} |

Both models assigned much greater normalized entropy to incorrect predictions than to correct ones. Softmax had a small but consistent error-AUROC advantage. DAEDL produced larger entropy values, while AUROC showed how well each score ranked incorrect predictions above correct ones.

DAEDL consistently assigned greater entropy to SVHN than to CIFAR-10. Softmax had higher OOD AUROC at seeds 3024 and 3026; DAEDL was higher at seed 3025. Neither method was consistently better across the three seeds.

![Figure 4 entropy distributions](../figures/final/figure4_entropy_distributions_seed3024.png)

## 7 Discussion

Softmax achieved higher clean classification accuracy in all three seeds. In this setup, DAEDL's uncertainty modelling did not improve classification performance. DAEDL consistently achieved lower ECE, whereas Softmax achieved lower NLL and Brier scores. These metrics can disagree because ECE measures binned confidence-frequency agreement, while NLL and Brier score evaluate the full probability distribution.

DAEDL produced higher entropy for correct, incorrect, ID, and OOD samples. That wider numerical range made mistakes and SVHN inputs visibly uncertain. AUROC, however, depends on how consistently the score ranks errors or distribution shifts, and DAEDL did not improve that ranking consistently.

Multiple seeds changed the interpretation of OOD performance. Seed 3024 favored Softmax, seed 3025 favored DAEDL, and seed 3026 favored Softmax. The three-seed mean favored Softmax, but the ordering was not consistent. With only three runs, the comparison is reported descriptively without a hypothesis test.

The density diagnostics show another trade-off. DAEDL embeddings remained rank deficient: global centered rank was 93, 99, and 101 out of 128, with 35, 29, and 27 constant dimensions. Per-class covariance ranks ranged from 76 to 95. Adaptive jitter remained around $10^{{-8}}$ to $10^{{-7}}$, and every density computation was finite. Density fitting remained numerically stable, although the reduced embedding rank is a limitation of this implementation.

This is an assignment-scale evaluation of DAEDL's central ideas. The paper used a different architecture and larger training setup. Here, the compact ResNet-18, 128-dimensional embedding, 30 epochs, and laptop hardware made matched multi-seed experiments feasible. The results apply to this local setup and are not an exact reproduction of the paper's broader benchmarks.

## 8 Limitations

The evaluation uses one in-distribution dataset, one OOD dataset, and three seeds. Corruption experiments were outside the final scope of this assignment. No exhaustive hyperparameter search was performed, and the study did not cover other architectures or alternative evidential methods. MPS and CPU adaptations introduce an engineering difference from the authors' CUDA-oriented environment. DAEDL's embedding space was rank deficient even though covariance stabilization remained finite. The findings apply to this controlled CIFAR-10/SVHN setup.

## 9 What I Learned

Before this assignment, I tended to read a classifier's probability as a direct statement of confidence. The experiment showed me that uncertainty needs separate evidence. A model can be accurate and still assign probabilities that do not match how often it is correct. It can also produce high entropy on difficult images without ranking every mistake or OOD sample better than another model.

I learned that calibration is not a single number. DAEDL had lower ECE in every seed, which initially looks like a clear calibration advantage. NLL and Brier score gave the opposite ranking. This result helped me understand that ECE depends on confidence bins, while proper scoring rules evaluate the full probability distribution and penalize confident errors differently. I should therefore explain which calibration property a metric measures instead of using “better calibrated” without qualification.

I also expected a method designed for uncertainty to outperform Softmax clearly on OOD detection. The result at seed 3024 initially pointed in the other direction. After running the other two seeds, my interpretation changed: DAEDL was better at seed 3025, while Softmax was better again at seed 3026. The full result was mixed across seeds.

The project also taught me to separate uncertainty magnitude from ranking quality. DAEDL produced much greater entropy on incorrect and SVHN inputs. Softmax still had slightly better error AUROC and a higher mean OOD AUROC. The separation between groups matters more than the absolute scale when models use different uncertainty parameterizations.

Codex wrote the code and ran the experiments. I found it especially useful for debugging and checking the experiment, not only for generating code. The workflow still needed tests for formulas and metrics, matched baselines, fixed checkpoint rules, checks for OOD score direction, probability normalization, and independent recomputation from saved predictions. It also helped me inspect failures, record changes, and keep mixed results in the final analysis.

Finally, I learned how theoretical ideas meet hardware and numerical constraints. MPS was suitable for neural-network training but could not run the float64 covariance calculations. Reloaded spectral-normalized checkpoints also required CPU evaluation in this environment. The embedding covariances were rank deficient, so the implementation needed adaptive diagonal loading and diagnostics. These hardware and numerical checks became part of the experimental workflow.

## 10 Conclusion

Codex implemented DAEDL and compared it with a matched Softmax baseline across three CIFAR-10 seeds. DAEDL consistently reduced ECE. Softmax achieved better accuracy, Macro F1, NLL, Brier score, and error-detection AUROC. Both methods assigned greater uncertainty to mistakes, and DAEDL assigned greater entropy to SVHN than CIFAR-10. The OOD ranking changed across seeds, and the added probabilistic structure did not improve every predictive or uncertainty measure.
{source_section}
{references_heading}

[1] Yoon, T., and Kim, H. (2024). “Uncertainty Estimation by Density Aware Evidential Deep Learning.” *Proceedings of the 41st International Conference on Machine Learning*, PMLR 235, 57217–57243. https://proceedings.mlr.press/v235/yoon24a.html

[2] Sensoy, M., Kaplan, L., and Kandemir, M. (2018). “Evidential Deep Learning to Quantify Classification Uncertainty.” *Advances in Neural Information Processing Systems 31*. https://papers.nips.cc/paper/2018/hash/a981f2b708044d6fb4a71a1463242520-Abstract.html

[3] Guo, C., Pleiss, G., Sun, Y., and Weinberger, K. Q. (2017). “On Calibration of Modern Neural Networks.” *Proceedings of the 34th International Conference on Machine Learning*, PMLR 70, 1321–1330. https://proceedings.mlr.press/v70/guo17a.html

[4] Krizhevsky, A. (2009). *Learning Multiple Layers of Features from Tiny Images*. Technical report, University of Toronto. https://www.cs.toronto.edu/~kriz/learning-features-2009-TR.pdf

[5] Netzer, Y., Wang, T., Coates, A., Bissacco, A., Wu, B., and Ng, A. Y. (2011). “Reading Digits in Natural Images with Unsupervised Feature Learning.” *NIPS Workshop on Deep Learning and Unsupervised Feature Learning*. https://cs.stanford.edu/~twangcat/papers/nips2011_housenumbers.pdf

[6] PyTorch Contributors. “MPS backend.” *PyTorch documentation*. https://docs.pytorch.org/docs/main/notes/mps.html

[7] Yoon, T., and Kim, H. “DAEDL official implementation.” GitHub. https://github.com/TaeseongYoon/DAEDL
"""


def build_docx(aggregate: pd.DataFrame, density: pd.DataFrame, source_url: str | None, output: Path) -> None:
    document = Document()
    configure_document(document)

    title = document.add_paragraph(style="Title")
    remove_paragraph_border(title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(120)
    title.add_run("Density-Aware Evidential Deep Learning for CIFAR-10\nClassification and Uncertainty Estimation")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(16)
    run = subtitle.add_run("CISC3024 Pattern Recognition\nAI Assignment #2")
    run.bold = True
    run.font.size = Pt(15)
    name = document.add_paragraph("Ziqi Huang")
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_before = Pt(18)
    date = document.add_paragraph("October 2026")
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_before = Pt(8)
    document.add_page_break()

    document.add_heading("1 Introduction", level=1)
    paragraph(document, "Image classifiers produce probability-like scores, but a high maximum score does not by itself show that the model knows when it may be wrong. This assignment used Codex to identify, implement, and evaluate a recent probability-based recognition method. The selected method was Density Aware Evidential Deep Learning (DAEDL), published at ICML 2024 [1].")
    paragraph(document, "The research question was: how does DAEDL compare with a conventional Softmax classifier in classification performance, calibration, prediction-error uncertainty, and out-of-distribution uncertainty in an assignment-scale CIFAR-10 experiment?")
    paragraph(document, f"The three-seed results show different strengths across the evaluation metrics. Softmax achieved higher accuracy ({value(aggregate, 'softmax', 'accuracy')} versus {value(aggregate, 'daedl', 'accuracy')}) and better NLL and Brier score. DAEDL achieved lower ECE in every seed ({value(aggregate, 'daedl', 'ece_15')} versus {value(aggregate, 'softmax', 'ece_15')}). Both methods assigned greater entropy to mistakes, and Softmax had a small error-AUROC advantage. The SVHN OOD ranking changed across seeds, so neither method was consistently better.")

    document.add_heading("2 AI-Assisted Algorithm Search and Selection", level=1)
    paragraph(document, "Codex searched for recent classification and recognition methods grounded in probability or Bayesian theory. Codex first compared DAEDL with conventional evidential deep learning. The candidates were screened by publication date, probabilistic basis, implementation availability, computational cost, and suitability for a CIFAR-10 comparison.")
    paragraph(document, "DAEDL was selected because it predicts a Dirichlet distribution, models density in learned feature space, and supports classification, calibration, error-detection, and OOD evaluation [1]. It was recent, directly relevant to image recognition, and small enough to adapt into a controlled CIFAR-10 study.")
    paragraph(document, "The search then checked the official PMLR record, the paper, and the authors' repository. Codex used the repository as a method reference and created a separate implementation adapted to this assignment [7]. The repository audit found CUDA assumptions, incomplete environment information, missing released checkpoints, and reproducibility gaps. Conventional evidential deep learning [2] and calibration work [3] supplied the background needed to interpret the method and metrics.")

    document.add_heading("3 Algorithm Description", level=1)
    document.add_heading("3.1 Softmax classification", level=2)
    paragraph(document, "A neural classifier maps an input image x to class logits. Softmax exponentiates and normalizes those logits to obtain one probability vector. The largest entry determines the predicted class. The operation guarantees non-negative probabilities summing to one, but it does not model a distribution over possible probability vectors.")
    add_equation(document, "equation1_probabilities.png", 6.1)
    paragraph(document, "The second expression above is the evidential counterpart. A Dirichlet random variable π represents uncertainty over categorical probability vectors, and its mean gives the predictive class probability. A larger total concentration S places more mass around that mean probability vector; a smaller value represents weaker evidence.")

    document.add_heading("3.2 Evidential training", level=2)
    paragraph(document, "Conventional EDL converts non-negative evidence into Dirichlet concentrations and trains a deterministic neural network to predict evidence [2]. DAEDL changes the training parameterization to an exponential of each logit. Once normalized by total strength, these concentrations produce exactly the Softmax probabilities during training. The loss combines expected squared classification error, predictive variance, and a KL regularizer that suppresses misleading evidence.")
    add_equation(document, "equation2_training.png", 6.65)
    paragraph(document, "The local implementation averages the complete per-example objective. This scales every component by the same batch-size factor, so the relative weight λ=0.05 is unchanged. Logits are clamped only at ±30 when exponentiated as an overflow guard.")

    document.add_heading("3.3 Feature density", level=2)
    paragraph(document, "After neural-network training, DAEDL extracts the 128-dimensional penultimate feature h(x) for every training image. A full-covariance Gaussian is fitted to each class, and the empirical class proportions form a mixture density. Training log-density extrema define a clipped score s(x) between zero and one.")
    add_equation(document, "equation3_density.png", 6.45)
    paragraph(document, "A test embedding near the fitted training features receives a larger density score. An embedding far from the fitted distribution receives a smaller score and therefore weaker evidence at prediction time. Because full covariance and log-density calculations are sensitive to numerical precision, the implementation performs this stage on CPU in float64 and adds scale-relative diagonal jitter when Cholesky factorization requires it.")

    document.add_heading("3.4 Density aware prediction", level=2)
    paragraph(document, "At inference, DAEDL multiplies each logit by the density score before constructing the Dirichlet concentrations. This is equivalent to an input-dependent temperature. As the score approaches zero, all concentrations approach one and the expected probabilities approach the uniform distribution.")
    add_equation(document, "equation4_inference.png", 6.25)
    paragraph(document, "This project used normalized predictive entropy to compare uncertainty across the two methods. The normalization places entropy in [0,1], so the same error-detection and OOD metric code can be used.")
    add_equation(document, "equation5_entropy.png", 3.2)

    add_caption(document, "Figure 1", "DAEDL inference pipeline used in this project. The diagram is original and was generated from the implemented method.")
    add_figure(document, FIGURES / "figure1_daedl_concept.png", 6.7)

    document.add_heading("4 AI-Assisted Implementation", level=1)
    document.add_heading("4.1 Project and environment setup", level=2)
    paragraph(document, "Codex created an isolated Python 3.11.5 environment with PyTorch 2.14.1 and torchvision 0.29.1 on an Apple M5 MacBook Air with 16 GB unified memory. The implementation covers data loading, model construction, Softmax and DAEDL training, density fitting, evaluation, artifact logging, and report figures.")
    paragraph(document, "Neural-network training used MPS. Selected checkpoints were evaluated on CPU, and DAEDL density estimation used CPU float64 because the MPS backend cannot store float64 tensors [6]. The pilot also found that a reloaded spectral-normalized checkpoint could produce non-finite outputs on the observed MPS evaluation path while the same weights were finite on CPU. The final protocol used the same CPU checkpoint-evaluation policy for both models.")

    document.add_heading("4.2 Data pipeline", level=2)
    paragraph(document, "Each CIFAR-10 example retained its original sample index so that splits and predictions could be audited. The 50,000 official training images were divided into a stratified 45,000-image training set and 5,000-image validation set. Training used random 32-by-32 crops with four-pixel padding, random horizontal flips, tensor conversion, and channel normalization with mean (0.4914, 0.4822, 0.4465) and standard deviation (0.2023, 0.1994, 0.2010). Validation, test, and density fitting used only tensor conversion and the same normalization.")
    paragraph(document, "The code saved the train, validation, and test indices for every run. Python, NumPy, PyTorch, and the data-loader generator were seeded. Within a seed, Softmax and DAEDL used identical indices; across seeds, the split and training randomness changed according to the recorded seed.")

    document.add_heading("4.3 Model implementation", level=2)
    paragraph(document, "Both methods used the same CIFAR-adapted ResNet-18. It has a 3-by-3 stride-one stem, base width 32, four residual stages, adaptive average pooling, a 128-dimensional ReLU embedding, and a ten-class linear classifier. The forward method exposes both logits and embeddings because DAEDL needs the embeddings after training.")
    paragraph(document, "Spectral normalization was applied to every convolution, learned shortcut projection, the embedding layer, and the classifier for both models. Sharing the feature extractor, data pipeline, and training budget keeps the comparison focused on the prediction and uncertainty mechanisms. Softmax uses cross-entropy. DAEDL uses its evidential objective, followed by feature-density fitting and density-aware inference.")

    document.add_heading("4.4 DAEDL density estimation", level=2)
    paragraph(document, "Density estimation begins after neural-network training. The selected checkpoint is loaded on CPU, and clean embeddings are extracted from all 45,000 training images. The implementation fits one full-covariance Gaussian per class and uses empirical class proportions as mixture weights. The minimum and maximum training log densities define the clipped score s(x) used at inference.")
    paragraph(document, "All covariance calculations use CPU float64. Each covariance is symmetrized, and scale-relative diagonal jitter with an absolute floor of 1e-8 is increased until Cholesky factorization succeeds. The run artifacts record global and per-class ranks, constant dimensions, jitter, the training density range, and non-finite fractions. The fitted GDA state and per-example training density values are also saved.")

    document.add_heading("4.5 Staged implementation workflow", level=2)
    paragraph(document, "The implementation was validated in stages. A two-epoch pilot used 5,000 training, 1,000 validation, and 1,000 test images to check finite losses and gradients, checkpoint reloads, GDA fitting, metrics, and artifact output. The first full experiment then trained both models on seed 3024 and evaluated classification, calibration, prediction errors, SVHN OOD detection, and density diagnostics. After those checks passed, the same configuration was then used for seeds 3025 and 3026.")

    document.add_heading("4.6 Validation testing and reproducibility", level=2)
    paragraph(document, "The 16 automated tests cover four groups. Probability and loss tests check positive concentrations, normalized probabilities, and finite DAEDL losses and gradients. Metric tests check NLL, Brier score, ECE, entropy, and AUROC/AUPR direction. Density tests cover shapes, CPU float64 state, rank-deficient covariance handling, and finite scores. Reproducibility tests cover stratified disjoint splits, exposed model outputs, and spectral-normalization coverage.")
    paragraph(document, "Every final run saved its configuration, environment, split indices, progress log, training history, best-validation and final checkpoints, per-example CIFAR-10 and SVHN predictions, metrics, runtime, and DAEDL density diagnostics where applicable. Aggregate tables were recomputed from the saved predictions. The six-run audit verified matched within-seed splits, checkpoint selection, probability normalization, finite outputs, positive DAEDL concentrations, score direction, density stability, and sample-SD aggregation.")

    document.add_heading("5 Experimental Setup", level=1)
    document.add_heading("5.1 Datasets and split", level=2)
    paragraph(document, "CIFAR-10 contains ten natural-object classes with 50,000 official training images and 10,000 official test images [4]. The training portion was divided into a stratified 45,000-image training set and 5,000-image validation set. Validation NLL selected checkpoints, while the official test set remained outside model selection.")
    paragraph(document, "SVHN supplied 10,000 deterministically selected test images for OOD evaluation [5]. CIFAR-10 contains natural objects; SVHN contains street-view house-number digits. Their visual domains and semantic classes differ, making SVHN a useful OOD dataset for this controlled comparison without implying that it is the only possible benchmark.")

    document.add_heading("5.2 Fair comparison design", level=2)
    paragraph(document, "Softmax and DAEDL shared the same backbone, spectral-normalization configuration, split, training augmentation, batch size, optimizer, epoch budget, learning-rate schedule, and checkpoint criterion within each seed. These controls reduce differences caused by the feature extractor or training resources. DAEDL necessarily differs in its evidential objective, post-training density model, and density-aware inference.")

    document.add_heading("5.3 Training protocol", level=2)
    paragraph(document, "Each run trained for 30 epochs with batch size 64, AdamW, an initial learning rate of 1e-3, weight decay 5e-4, and cosine decay to 1e-5. Both models selected the checkpoint with minimum validation NLL. This provides a common probability-based criterion even though their training objectives differ. The best-validation checkpoint and the final-epoch checkpoint were both saved, and the test set was never used for selection.")
    add_caption(document, "Table 1", "Experimental configuration.")
    add_table(document, ["Setting", "Value"], [
        ["Models", "Softmax and DAEDL"],
        ["Backbone", "Spectrally normalized CIFAR ResNet-18, base width 32"],
        ["Embedding", "128 dimensions"],
        ["Split", "45,000 train / 5,000 validation / 10,000 test"],
        ["Seeds", "3024, 3025, 3026"],
        ["Optimizer", "AdamW; initial LR 1e-3; weight decay 5e-4"],
        ["Schedule", "Cosine decay to 1e-5 over 30 epochs"],
        ["Batch size", "64"],
        ["Selection", "Minimum validation NLL"],
        ["Hardware", "Apple M5 MacBook Air; 16 GB unified memory"],
        ["Devices", "MPS training; CPU evaluation and float64 GDA"],
        ["OOD protocol", "CIFAR-10 ID; SVHN OOD; normalized entropy; OOD positive"],
    ], [1.55, 5.0])

    document.add_heading("5.4 Three-seed replication", level=2)
    paragraph(document, "Seeds 3024, 3025, and 3026 each defined a complete model run. The seed controlled the stratified train-validation split, Python and NumPy random states, PyTorch initialization and training randomness, and the data-loader generator. Softmax and DAEDL shared indices within each seed. Results are reported as mean ± sample SD across the three runs, not across individual images. Three seeds were used to check whether the comparison remained consistent under different data splits and training randomness.")

    document.add_heading("5.5 Classification evaluation", level=2)
    paragraph(document, "Accuracy is the fraction of correct CIFAR-10 predictions. Macro F1 averages class-level F1 scores with equal weight for each class. CIFAR-10 is balanced, but Macro F1 still checks whether performance is distributed across classes instead of relying only on total accuracy.")

    document.add_heading("5.6 Calibration evaluation", level=2)
    paragraph(document, "Negative log-likelihood uses the probability assigned to the true class and gives a large penalty to confident mistakes. ECE divides predictions into 15 equal-width confidence bins and computes the count-weighted difference between mean confidence and observed accuracy in each bin. The multiclass Brier score is the mean squared difference between the predicted probability vector and the one-hot target. Using all three metrics gives a broader view of probability quality than relying on a single calibration measure.")

    document.add_heading("5.7 Prediction-error uncertainty", level=2)
    paragraph(document, "This experiment asks whether uncertainty can identify predictions that are likely to be wrong. Correct and incorrect CIFAR-10 predictions are separated using the ground-truth labels. Normalized predictive entropy lies in [0,1]; larger values indicate a more diffuse class distribution. Prediction error is the positive class. AUROC measures ranking across thresholds, and AUPR summarizes precision-recall performance for the error class.")

    document.add_heading("5.8 OOD evaluation", level=2)
    paragraph(document, "The OOD experiment asks whether the models assign greater uncertainty to inputs outside the CIFAR-10 distribution. CIFAR-10 is labelled ID, SVHN is labelled OOD, and OOD is the positive class. Normalized predictive entropy is the common score for both models, which avoids comparing method-specific scales. DAEDL's secondary score, 1-s(x), is evaluated separately as density-derived uncertainty.")

    document.add_heading("5.9 Numerical stability and device policy", level=2)
    paragraph(document, "Training used MPS, while loaded-checkpoint evaluation and feature extraction used CPU for both models. DAEDL covariance and density calculations used CPU float64, adaptive diagonal jitter, and Cholesky checks. Evaluation also checked finite outputs, probability sums, positive concentrations, density bounds, and score orientation before aggregation.")

    document.add_heading("5.10 Experiment questions", level=2)
    add_caption(document, "Table 2", "Evaluation questions and reported metrics.")
    add_table(document, ["Experiment", "Main question", "Metrics"], [
        ["Classification", "How accurately does the model classify CIFAR-10?", "Accuracy; Macro F1"],
        ["Calibration", "Do predicted probabilities match empirical correctness?", "NLL; ECE; Brier"],
        ["Error uncertainty", "Can uncertainty identify incorrect predictions?", "Entropy; AUROC; AUPR"],
        ["OOD uncertainty", "Can uncertainty distinguish CIFAR-10 from SVHN?", "Entropy; AUROC; AUPR"],
    ], [1.35, 3.75, 1.45])

    document.add_page_break()
    document.add_heading("6 Results", level=1)
    document.add_heading("6.1 Classification and calibration", level=2)
    add_caption(document, "Table 3", "Clean CIFAR-10 classification and calibration. Values are mean ± sample SD across three training seeds.")
    add_table(document, ["Metric", "Softmax", "DAEDL", "Preferred direction"], [
        ["Accuracy", value(aggregate, "softmax", "accuracy"), value(aggregate, "daedl", "accuracy"), "Higher"],
        ["Macro F1", value(aggregate, "softmax", "macro_f1"), value(aggregate, "daedl", "macro_f1"), "Higher"],
        ["NLL", value(aggregate, "softmax", "nll"), value(aggregate, "daedl", "nll"), "Lower"],
        ["ECE-15", value(aggregate, "softmax", "ece_15"), value(aggregate, "daedl", "ece_15"), "Lower"],
        ["Brier score", value(aggregate, "softmax", "brier"), value(aggregate, "daedl", "brier"), "Lower"],
    ], [1.65, 1.65, 1.65, 1.55])
    paragraph(document, "Softmax was consistently about one percentage point more accurate. Its accuracy exceeded DAEDL by 0.0099, 0.0130, and 0.0105 at seeds 3024, 3025, and 3026. Macro F1 followed the same pattern. Both methods nevertheless reached close to 89–90% accuracy in the compact 30-epoch setup.")
    paragraph(document, "DAEDL had lower ECE in every seed, but NLL and Brier score favored Softmax in every seed. The calibration metrics therefore do not all support the same ranking. ECE summarizes binned agreement between confidence and observed accuracy. NLL and Brier score are proper scoring rules that evaluate the complete predictive distribution and respond differently to probability sharpness and errors.")
    add_caption(document, "Figure 2", "Three-seed headline comparison. Points show individual seeds; black bars show mean ± sample SD.")
    add_figure(document, FIGURES / "figure2_three_seed_headline.png", 6.6)

    document.add_heading("6.2 Reliability", level=2)
    paragraph(document, "The seed-3024 reliability diagram illustrates why a single calibration summary needs context. DAEDL tracks the ideal diagonal more closely in several high-confidence bins, consistent with its lower ECE. The plot shows seed 3024 as an example. Bins with fewer than 50 test examples are suppressed for readability, while ECE is calculated from all 15 equal-width bins.")
    add_caption(document, "Figure 3", "CIFAR-10 reliability at the predefined representative seed 3024.")
    add_figure(document, FIGURES / "figure3_reliability_seed3024.png", 4.7)

    document.add_heading("6.3 Prediction error uncertainty", level=2)
    paragraph(document, f"Both methods assigned substantially greater normalized entropy to incorrect predictions. Across seeds, Softmax's correct and incorrect mean entropies were {value(aggregate, 'softmax', 'error_correct_mean_entropy')} and {value(aggregate, 'softmax', 'error_incorrect_mean_entropy')}; DAEDL's were {value(aggregate, 'daedl', 'error_correct_mean_entropy')} and {value(aggregate, 'daedl', 'error_incorrect_mean_entropy')}. The distributions therefore contain useful information about mistakes for both methods.")
    paragraph(document, "DAEDL produced larger entropy values. AUROC instead measures how well uncertainty ranks incorrect predictions above correct ones. Both models separated incorrect from correct predictions well, with Softmax achieving a slightly higher AUROC in every seed. AUPR changed paired direction.")

    document.add_heading("6.4 SVHN Out-of-Distribution Uncertainty", level=2)
    add_caption(document, "Table 4", "Error and SVHN OOD detection using normalized predictive entropy. Values are mean ± sample SD across three seeds.")
    add_table(document, ["Metric", "Softmax", "DAEDL"], [
        ["Error AUROC", value(aggregate, "softmax", "error_auroc"), value(aggregate, "daedl", "error_auroc")],
        ["Error AUPR", value(aggregate, "softmax", "error_aupr"), value(aggregate, "daedl", "error_aupr")],
        ["SVHN OOD AUROC", value(aggregate, "softmax", "ood_auroc"), value(aggregate, "daedl", "ood_auroc")],
        ["SVHN OOD AUPR", value(aggregate, "softmax", "ood_aupr"), value(aggregate, "daedl", "ood_aupr")],
    ], [2.5, 1.9, 1.9])
    paragraph(document, f"DAEDL consistently assigned greater entropy to SVHN than to CIFAR-10: its ID and OOD mean entropies were {value(aggregate, 'daedl', 'ood_id_mean_entropy')} and {value(aggregate, 'daedl', 'ood_mean_entropy')}. Softmax showed the same direction at a lower absolute scale. Softmax had higher OOD AUROC at seeds 3024 and 3026; DAEDL was higher at seed 3025. Neither method was consistently better across the three seeds.")
    paragraph(document, "DAEDL's secondary density-only uncertainty had AUROC values of 0.8641, 0.8420, and 0.8498 at seeds 3024–3026. The density score exceeded DAEDL entropy only at seed 3024, so the improvement was not consistent across seeds.")
    add_caption(document, "Figure 4", "Normalized predictive-entropy distributions at seed 3024. Numeric conclusions use all three seeds.")
    add_figure(document, FIGURES / "figure4_entropy_distributions_seed3024.png", 6.35)

    document.add_heading("7 Discussion", level=1)
    document.add_heading("7.1 Accuracy and uncertainty", level=2)
    paragraph(document, "Softmax achieved higher clean classification accuracy in all three seeds. In this setup, DAEDL's uncertainty modelling did not improve classification performance.")
    document.add_heading("7.2 Calibration is Multidimensional", level=2)
    paragraph(document, "DAEDL consistently achieved lower ECE, whereas Softmax achieved lower NLL and Brier scores. These metrics can disagree because ECE measures binned confidence-frequency agreement, while NLL and Brier score evaluate the full probability distribution and penalize probability errors differently.")
    document.add_heading("7.3 Magnitude and usefulness", level=2)
    paragraph(document, "DAEDL produced higher entropy for correct, incorrect, ID, and OOD samples. That wider numerical range made mistakes and SVHN inputs visibly uncertain. AUROC, however, depends on how consistently the score ranks risky inputs, and DAEDL did not improve that ranking consistently.")
    document.add_heading("7.4 Seed dependence", level=2)
    paragraph(document, "Multiple seeds changed the OOD interpretation. Seed 3024 favored Softmax, seed 3025 favored DAEDL, and seed 3026 favored Softmax. The three-seed mean favored Softmax, but the ordering was not consistent. With only three runs, the comparison is reported descriptively without a hypothesis test.")
    document.add_heading("7.5 Assignment scale adaptation", level=2)
    paragraph(document, "The original DAEDL work used a different architecture and larger training setup [1]. This project used a compact ResNet-18, 128-dimensional embeddings, 30 epochs, and laptop hardware so that a matched three-seed comparison was feasible. It evaluates the method's central ideas at assignment scale; the results are specific to this setup.")

    document.add_heading("8 Limitations", level=1)
    paragraph(document, "The evaluation uses one in-distribution dataset, one OOD dataset, and three seeds. No exhaustive hyperparameter search was performed, and the study did not cover other architectures or evidential baselines. MPS and CPU adaptations differ from the authors' CUDA-oriented environment. The DAEDL feature space was rank deficient even though covariance stabilization remained finite.")
    paragraph(document, f"Global centered embedding rank was {int(density.iloc[0].global_embedding_rank)}, {int(density.iloc[1].global_embedding_rank)}, and {int(density.iloc[2].global_embedding_rank)} out of 128, while constant dimensions were {int(density.iloc[0].constant_dimensions)}, {int(density.iloc[1].constant_dimensions)}, and {int(density.iloc[2].constant_dimensions)}. Per-class covariance ranks spanned {int(density.class_rank_min.min())}–{int(density.class_rank_max.max())}. Jitter remained around 1e-8 to 1e-7 and every non-finite fraction was zero. Density fitting remained numerically stable, although the reduced embedding rank is a limitation of this implementation.")
    paragraph(document, "Corruption experiments were outside the final scope of this assignment.")

    document.add_page_break()
    document.add_heading("9 What I Learned", level=1)
    paragraph(document, "Before this assignment, I tended to read a classifier's probability as a direct statement of confidence. The experiment showed me that uncertainty needs separate evidence. A model can be accurate and still assign probabilities that do not match how often it is correct.")
    paragraph(document, "I learned that calibration is not a single number. DAEDL had lower ECE in every seed, which initially looks like a clear calibration advantage. NLL and Brier score gave the opposite ranking. This result helped me understand that ECE depends on confidence bins, while proper scoring rules evaluate the full probability distribution and penalize confident errors differently. I should therefore explain which calibration property a metric measures instead of using ‘better calibrated’ without qualification.")
    paragraph(document, "I also expected a method designed for uncertainty to outperform Softmax clearly on OOD detection. The result at seed 3024 initially pointed in the other direction. After running the other two seeds, my interpretation changed: DAEDL was better at seed 3025, while Softmax was better again at seed 3026. The full result was mixed across seeds.")
    paragraph(document, "The project also taught me to separate uncertainty magnitude from ranking quality. DAEDL produced much greater entropy on incorrect and SVHN inputs. Softmax still had slightly better error AUROC and a higher mean OOD AUROC. The separation between groups matters more than the absolute scale when models use different uncertainty parameterizations.")
    paragraph(document, "Codex wrote the code and ran the experiments. I found it especially useful for debugging and checking the experiment, not only for generating code. The workflow still needed tests for formulas and metrics, matched baselines, fixed checkpoint rules, checks for OOD score direction, probability normalization, and independent recomputation from saved predictions. It also helped me inspect failures, record changes, and keep mixed results in the final analysis.")
    paragraph(document, "Finally, I learned how theory meets hardware and numerical constraints. MPS could train the networks but could not run float64 covariance calculations, and reloaded spectral-normalized checkpoints required CPU evaluation in this environment. Rank-deficient embedding covariances required adaptive diagonal loading and diagnostics. These hardware and numerical checks became part of the experimental workflow.")

    document.add_heading("10 Conclusion", level=1)
    paragraph(document, "Codex implemented DAEDL and compared it with a matched Softmax baseline across three CIFAR-10 seeds. DAEDL consistently reduced ECE. Softmax achieved better accuracy, Macro F1, NLL, Brier score, and error-detection AUROC. Both methods assigned greater uncertainty to mistakes, and DAEDL assigned greater entropy to SVHN than CIFAR-10. The OOD ranking changed across seeds, and the added probabilistic structure did not improve every predictive or uncertainty measure.")

    if source_url:
        document.add_heading("11 Source Code", level=1)
        paragraph(document, "The complete source code, experiment scripts, configurations, and reproducibility files for this assignment are available at:")
        p = document.add_paragraph()
        add_hyperlink(p, source_url, source_url)

    document.add_page_break()
    document.add_heading("12 References" if source_url else "References", level=1)
    references = [
        ("[1] Yoon, T., and Kim, H. (2024). “Uncertainty Estimation by Density Aware Evidential Deep Learning.” Proceedings of the 41st International Conference on Machine Learning, PMLR 235, 57217–57243. ", "https://proceedings.mlr.press/v235/yoon24a.html"),
        ("[2] Sensoy, M., Kaplan, L., and Kandemir, M. (2018). “Evidential Deep Learning to Quantify Classification Uncertainty.” Advances in Neural Information Processing Systems 31. ", "https://papers.nips.cc/paper/2018/hash/a981f2b708044d6fb4a71a1463242520-Abstract.html"),
        ("[3] Guo, C., Pleiss, G., Sun, Y., and Weinberger, K. Q. (2017). “On Calibration of Modern Neural Networks.” Proceedings of the 34th International Conference on Machine Learning, PMLR 70, 1321–1330. ", "https://proceedings.mlr.press/v70/guo17a.html"),
        ("[4] Krizhevsky, A. (2009). Learning Multiple Layers of Features from Tiny Images. Technical report, University of Toronto. ", "https://www.cs.toronto.edu/~kriz/learning-features-2009-TR.pdf"),
        ("[5] Netzer, Y., Wang, T., Coates, A., Bissacco, A., Wu, B., and Ng, A. Y. (2011). “Reading Digits in Natural Images with Unsupervised Feature Learning.” NIPS Workshop on Deep Learning and Unsupervised Feature Learning. ", "https://cs.stanford.edu/~twangcat/papers/nips2011_housenumbers.pdf"),
        ("[6] PyTorch Contributors. “MPS backend.” PyTorch documentation. ", "https://docs.pytorch.org/docs/main/notes/mps.html"),
        ("[7] Yoon, T., and Kim, H. “DAEDL official implementation.” GitHub. ", "https://github.com/TaeseongYoon/DAEDL"),
    ]
    for text, url in references:
        p = document.add_paragraph(style="References")
        p.add_run(text)
        add_hyperlink(p, url, url)

    properties = document.core_properties
    properties.title = "Density-Aware Evidential Deep Learning for CIFAR-10 Classification and Uncertainty Estimation"
    properties.subject = "CISC3024 Pattern Recognition AI Assignment #2"
    properties.author = "Ziqi Huang"
    properties.keywords = "DAEDL, CIFAR-10, uncertainty, calibration, evidential deep learning"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-url", help="Verified public source-code URL; omit until publication")
    parser.add_argument("--docx", type=Path, default=REPORT / "CISC3024_AI_Assignment2_Final.docx")
    parser.add_argument("--markdown", type=Path, default=REPORT / "CISC3024_AI_Assignment2_Final.md")
    args = parser.parse_args()
    aggregate, _, density = load_results()
    args.markdown.write_text(build_markdown(aggregate, density, args.source_url))
    build_docx(aggregate, density, args.source_url, args.docx)
    print(json.dumps({"docx": str(args.docx), "markdown": str(args.markdown), "source_url": args.source_url}, indent=2))


if __name__ == "__main__":
    main()
